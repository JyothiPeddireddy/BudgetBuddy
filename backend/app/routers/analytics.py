from io import BytesIO
from datetime import date, datetime
from calendar import monthrange
from typing import List, Optional

from fastapi import APIRouter, Depends, Query, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func
from pydantic import BaseModel

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.database import get_db
from app.models.expense import Expense
from app.models.income import Income
from app.models.savings_goal import SavingsGoal
from app.core.deps import get_current_user
from app.core.permissions import (
    require_premium_plus,
    require_export_access,
    has_min_role,
    UserRole,
)

from app.routers.reports import get_month_data, get_accounts_data

from app.crud.savings_goal import get_goal_contribution_trend  # <-- NEW

router = APIRouter(
    prefix="/analytics",
    tags=["Analytics"],
)


# ============================================================
# RESPONSE SCHEMAS
# ============================================================

class CategorySpending(BaseModel):
    category: str
    total: float


class MonthlyTrendPoint(BaseModel):
    month: str
    label: str
    total_income: float
    total_expenses: float


class GoalProgress(BaseModel):
    id: int
    title: str
    target_amount: float
    current_amount: float
    progress_percent: int
    status: str
    target_date: Optional[date] = None
    icon: str = "laptop"


class AnalyticsSummary(BaseModel):
    month: int
    year: int
    total_income: float
    total_expenses: float
    net_savings: float
    savings_rate: float
    total_balance: float


# ============================================================
# HELPERS
# ============================================================

def months_back_start(months: int) -> date:
    today = date.today()
    total_index = today.year * 12 + (today.month - 1) - (months - 1)
    start_year = total_index // 12
    start_month = total_index % 12 + 1
    return date(start_year, start_month, 1)


def month_sequence(start: date, end: date):
    year, month = start.year, start.month
    while (year, month) <= (end.year, end.month):
        yield (year, month)
        month += 1
        if month > 12:
            month = 1
            year += 1


MONTH_LABELS = [
    "Jan", "Feb", "Mar", "Apr", "May", "Jun",
    "Jul", "Aug", "Sep", "Oct", "Nov", "Dec",
]


def get_category_spending(month, year, db: Session, current_user):
    query = db.query(
        Expense.category,
        func.sum(Expense.amount).label("total"),
    ).filter(Expense.user_id == current_user.id)

    if month and year:
        start_date = date(year, month, 1)
        end_date = date(year, month, monthrange(year, month)[1])
        query = query.filter(Expense.date >= start_date, Expense.date <= end_date)

    rows = (
        query.group_by(Expense.category)
        .order_by(func.sum(Expense.amount).desc())
        .all()
    )

    return [{"category": category, "total": float(total or 0)} for category, total in rows]


def get_monthly_trend_points(months: int, db: Session, current_user):
    today = date.today()
    start_date = months_back_start(months)

    expenses = (
        db.query(Expense)
        .filter(
            Expense.user_id == current_user.id,
            Expense.date >= start_date,
            Expense.date <= today,
        )
        .all()
    )

    incomes = (
        db.query(Income)
        .filter(
            Income.user_id == current_user.id,
            Income.date >= start_date,
            Income.date <= today,
        )
        .all()
    )

    expense_totals = {}
    for e in expenses:
        key = (e.date.year, e.date.month)
        expense_totals[key] = expense_totals.get(key, 0) + float(e.amount or 0)

    income_totals = {}
    for i in incomes:
        key = (i.date.year, i.date.month)
        income_totals[key] = income_totals.get(key, 0) + float(i.amount or 0)

    points = []
    for year, month in month_sequence(start_date, today):
        key = (year, month)
        points.append({
            "month": f"{year:04d}-{month:02d}",
            "label": f"{MONTH_LABELS[month - 1]} {year}",
            "total_income": income_totals.get(key, 0),
            "total_expenses": expense_totals.get(key, 0),
        })

    return points


def get_savings_progress_rows(db: Session, current_user):
    """Returns all savings goals with their completion percentage."""
    goals = (
        db.query(SavingsGoal)
        .filter(SavingsGoal.user_id == current_user.id)
        .order_by(SavingsGoal.created_at.desc())
        .all()
    )

    results = []
    for g in goals:
        target = float(g.target_amount or 0)
        current = float(g.current_amount or 0)
        progress = round((current / target) * 100) if target > 0 else 0
        results.append({
            "id": g.id,
            "title": g.title,
            "target_amount": target,
            "current_amount": current,
            "progress_percent": min(progress, 100),
            "status": (g.status or "in_progress"),
            "target_date": g.target_date,
            "icon": g.icon,
        })

    return results


def get_category_spending_range(start_date: date, end_date: date, db: Session, current_user):
    """Same as get_category_spending but filtered by an explicit date range
    instead of a single month. Used by Premium's custom range picker."""
    rows = (
        db.query(Expense.category, func.sum(Expense.amount).label("total"))
        .filter(
            Expense.user_id == current_user.id,
            Expense.date >= start_date,
            Expense.date <= end_date,
        )
        .group_by(Expense.category)
        .order_by(func.sum(Expense.amount).desc())
        .all()
    )
    return [{"category": category, "total": float(total or 0)} for category, total in rows]


def get_summary_range(start_date: date, end_date: date, db: Session, current_user):
    """Income/expense/net-savings totals over an arbitrary date range
    (rather than a single calendar month)."""
    total_income = (
        db.query(func.sum(Income.amount))
        .filter(
            Income.user_id == current_user.id,
            Income.date >= start_date,
            Income.date <= end_date,
        )
        .scalar()
        or 0
    )
    total_expenses = (
        db.query(func.sum(Expense.amount))
        .filter(
            Expense.user_id == current_user.id,
            Expense.date >= start_date,
            Expense.date <= end_date,
        )
        .scalar()
        or 0
    )
    total_income = float(total_income)
    total_expenses = float(total_expenses)
    net_savings = total_income - total_expenses
    savings_rate = round((net_savings / total_income) * 100, 1) if total_income > 0 else 0.0

    _accounts, total_balance = get_accounts_data(db, current_user)

    return {
        "start_date": start_date,
        "end_date": end_date,
        "total_income": total_income,
        "total_expenses": total_expenses,
        "net_savings": net_savings,
        "savings_rate": savings_rate,
        "total_balance": total_balance,
    }


def get_category_trend_rows(months: int, db: Session, current_user):
    """Per-category spending broken out by month, for the stacked bar chart.
    Returns [{month, label, categories: {category: total, ...}}, ...]."""
    today = date.today()
    start_date = months_back_start(months)

    expenses = (
        db.query(Expense)
        .filter(
            Expense.user_id == current_user.id,
            Expense.date >= start_date,
            Expense.date <= today,
        )
        .all()
    )

    # {(year, month): {category: total}}
    grouped = {}
    for e in expenses:
        key = (e.date.year, e.date.month)
        grouped.setdefault(key, {})
        grouped[key][e.category] = grouped[key].get(e.category, 0) + float(e.amount or 0)

    points = []
    for year, month in month_sequence(start_date, today):
        key = (year, month)
        points.append({
            "month": f"{year:04d}-{month:02d}",
            "label": f"{MONTH_LABELS[month - 1]} {year}",
            "categories": grouped.get(key, {}),
        })

    return points


def get_analytics_summary_data(month, year, db: Session, current_user):
    today = date.today()
    target_month = month or today.month
    target_year = year or today.year

    (
        _incomes,
        _expenses,
        total_income,
        total_expenses,
        net_savings,
        _report_date,
    ) = get_month_data(target_month, target_year, db, current_user)

    _accounts, total_balance = get_accounts_data(db, current_user)

    savings_rate = round((net_savings / total_income) * 100, 1) if total_income > 0 else 0.0

    return {
        "month": target_month,
        "year": target_year,
        "total_income": total_income,
        "total_expenses": total_expenses,
        "net_savings": net_savings,
        "savings_rate": savings_rate,
        "total_balance": total_balance,
    }


# ============================================================
# GET /analytics/spending-by-category
# Open to all roles (User tier gets "basic charts" per your table).
# ============================================================

@router.get("/spending-by-category", response_model=List[CategorySpending])
def spending_by_category(
    month: Optional[int] = Query(None, ge=1, le=12),
    year: Optional[int] = Query(None, ge=2000, le=2100),
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user),
):
    # Custom date range is a Premium+ feature. If a User-tier account sends
    # start_date/end_date anyway (e.g. by calling the API directly), ignore
    # them silently and fall back to the month/year (or current-month) view
    # rather than erroring — keeps this endpoint usable by every tier.
    if start_date and end_date and has_min_role(current_user, UserRole.PREMIUM):
        rows = get_category_spending_range(start_date, end_date, db, current_user)
    else:
        rows = get_category_spending(month, year, db, current_user)
    return [CategorySpending(**row) for row in rows]


# ============================================================
# GET /analytics/monthly-trend
# CHANGED: Premium+ only. This is your "6-12 months history" feature,
# which the table only grants to Premium and Admin.
# ============================================================

@router.get("/monthly-trend", response_model=List[MonthlyTrendPoint])
def monthly_trend(
    months: int = Query(6, ge=1, le=12),
    db: Session = Depends(get_db),
    current_user=Depends(require_premium_plus),  # was: get_current_user
):
    points = get_monthly_trend_points(months, db, current_user)
    return [MonthlyTrendPoint(**p) for p in points]


# ============================================================
# GET /analytics/savings-progress
# Open to all roles (savings goal progress % is shown to every tier).
# ============================================================

@router.get("/savings-progress", response_model=List[GoalProgress])
def savings_progress(
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user),
):
    rows = get_savings_progress_rows(db, current_user)
    return [GoalProgress(**row) for row in rows]


# ============================================================
# GET /analytics/goals/{goal_id}/contribution-trend
# Premium+ only. Contribution history for one goal, oldest-first,
# with a running cumulative total.
# ============================================================

class ContributionPoint(BaseModel):
    date: str
    amount: float
    cumulative: float


@router.get("/goals/{goal_id}/contribution-trend", response_model=List[ContributionPoint])
def goal_contribution_trend(
    goal_id: int,
    db: Session = Depends(get_db),
    current_user=Depends(require_premium_plus),
):
    points = get_goal_contribution_trend(db, goal_id, current_user.id)
    if points is None:
        raise HTTPException(status_code=404, detail="Savings goal not found")
    return [ContributionPoint(**p) for p in points]


# ============================================================
# GET /analytics/summary
# Open to all roles (used for both basic and full dashboards).
# ============================================================

@router.get("/summary", response_model=AnalyticsSummary)
def analytics_summary(
    month: Optional[int] = Query(None, ge=1, le=12),
    year: Optional[int] = Query(None, ge=2000, le=2100),
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user),
):
    if start_date and end_date and has_min_role(current_user, UserRole.PREMIUM):
        # AnalyticsSummary expects month/year fields -- fill them with the
        # range start for display purposes; the totals themselves are the
        # real range-based numbers.
        range_data = get_summary_range(start_date, end_date, db, current_user)
        data = {
            "month": start_date.month,
            "year": start_date.year,
            "total_income": range_data["total_income"],
            "total_expenses": range_data["total_expenses"],
            "net_savings": range_data["net_savings"],
            "savings_rate": range_data["savings_rate"],
            "total_balance": range_data["total_balance"],
        }
    else:
        data = get_analytics_summary_data(month, year, db, current_user)
    return AnalyticsSummary(**data)


# ============================================================
# GET /analytics/comparison
# Premium+ only. This month vs last month, % change.
# ============================================================

class ComparisonMetric(BaseModel):
    current: float
    previous: float
    change_percent: Optional[float] = None  # None when previous is 0 (can't divide)


class AnalyticsComparison(BaseModel):
    current_month: str   # "2026-09"
    previous_month: str  # "2026-08"
    income: ComparisonMetric
    expenses: ComparisonMetric
    net_savings: ComparisonMetric


def _pct_change(current: float, previous: float) -> Optional[float]:
    if previous == 0:
        return None
    return round(((current - previous) / abs(previous)) * 100, 1)


@router.get("/comparison", response_model=AnalyticsComparison)
def analytics_comparison(
    month: Optional[int] = Query(None, ge=1, le=12),
    year: Optional[int] = Query(None, ge=2000, le=2100),
    db: Session = Depends(get_db),
    current_user=Depends(require_premium_plus),
):
    today = date.today()
    target_month = month or today.month
    target_year = year or today.year

    # Previous month, handling January -> December of prior year
    if target_month == 1:
        prev_month, prev_year = 12, target_year - 1
    else:
        prev_month, prev_year = target_month - 1, target_year

    current_data = get_analytics_summary_data(target_month, target_year, db, current_user)
    previous_data = get_analytics_summary_data(prev_month, prev_year, db, current_user)

    return AnalyticsComparison(
        current_month=f"{target_year:04d}-{target_month:02d}",
        previous_month=f"{prev_year:04d}-{prev_month:02d}",
        income=ComparisonMetric(
            current=current_data["total_income"],
            previous=previous_data["total_income"],
            change_percent=_pct_change(current_data["total_income"], previous_data["total_income"]),
        ),
        expenses=ComparisonMetric(
            current=current_data["total_expenses"],
            previous=previous_data["total_expenses"],
            change_percent=_pct_change(current_data["total_expenses"], previous_data["total_expenses"]),
        ),
        net_savings=ComparisonMetric(
            current=current_data["net_savings"],
            previous=previous_data["net_savings"],
            change_percent=_pct_change(current_data["net_savings"], previous_data["net_savings"]),
        ),
    )


# ============================================================
# GET /analytics/category-trend
# Premium+ only. Per-category spending across months (stacked bar).
# ============================================================

class CategoryTrendPoint(BaseModel):
    month: str
    label: str
    categories: dict[str, float]


@router.get("/category-trend", response_model=List[CategoryTrendPoint])
def category_trend(
    months: int = Query(6, ge=1, le=12),
    db: Session = Depends(get_db),
    current_user=Depends(require_premium_plus),
):
    points = get_category_trend_rows(months, db, current_user)
    return [CategoryTrendPoint(**p) for p in points]


# ============================================================
# EXPORT — shared styles
# ============================================================

def _pdf_styles():
    styles = getSampleStyleSheet()

    return {
        "title": ParagraphStyle(
            "ReportTitle", parent=styles["Title"], fontSize=20, leading=24,
            alignment=TA_CENTER, textColor=colors.HexColor("#071A3D"), spaceAfter=5,
        ),
        "subtitle": ParagraphStyle(
            "ReportSubtitle", parent=styles["Normal"], fontSize=11,
            alignment=TA_CENTER, textColor=colors.HexColor("#64748B"), spaceAfter=8,
        ),
        "month": ParagraphStyle(
            "ReportMonth", parent=styles["Normal"], fontSize=13,
            alignment=TA_CENTER, textColor=colors.HexColor("#10B981"), spaceAfter=4,
        ),
        "generated": ParagraphStyle(
            "GeneratedAt", parent=styles["Normal"], fontSize=9,
            alignment=TA_CENTER, textColor=colors.HexColor("#94A3B8"), spaceAfter=20,
        ),
        "heading": ParagraphStyle(
            "ReportHeading", parent=styles["Heading2"], fontSize=13, leading=16,
            textColor=colors.HexColor("#071A3D"), spaceBefore=16, spaceAfter=8,
        ),
        "empty": ParagraphStyle(
            "EmptySection", parent=styles["Normal"], fontSize=9.5,
            textColor=colors.HexColor("#64748B"), spaceAfter=4,
        ),
        "footer": ParagraphStyle(
            "Footer", parent=styles["Normal"], fontSize=8,
            alignment=TA_CENTER, textColor=colors.HexColor("#64748B"),
        ),
    }


# ============================================================
# GET /analytics/export/pdf
# CHANGED: Premium+ only. User tier's "Limited" export doesn't include
# PDF per your table — build a separate CSV-only export for User tier
# if/when you want that (not included here, see note below).
# ============================================================

@router.get("/export/pdf")
def export_analytics_pdf(
    month: int = Query(..., ge=1, le=12),
    year: int = Query(..., ge=2000, le=2100),
    months: int = Query(6, ge=1, le=12),
    db: Session = Depends(get_db),
    current_user=Depends(require_export_access),  # was: get_current_user
):
    summary = get_analytics_summary_data(month, year, db, current_user)
    category_rows = get_category_spending(month, year, db, current_user)
    trend_points = get_monthly_trend_points(months, db, current_user)
    goal_rows = get_savings_progress_rows(db, current_user)

    month_name = date(year, month, 1).strftime("%B %Y")
    generated_at = datetime.now().strftime("%d %b %Y, %I:%M %p")

    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer, pagesize=A4,
        leftMargin=40, rightMargin=40, topMargin=40, bottomMargin=40,
    )

    s = _pdf_styles()
    story = []

    story.append(Paragraph("BUDGET BUDDY", s["title"]))
    story.append(Paragraph("ANALYTICS REPORT", s["subtitle"]))
    story.append(Paragraph(month_name, s["month"]))
    story.append(Paragraph(f"Generated on {generated_at}", s["generated"]))

    story.append(Paragraph("Financial Summary", s["heading"]))

    summary_rows = [
        ["Total Income", "Total Expenses", "Net Savings", "Savings Rate", "Total Balance"],
        [
            f"Rs. {summary['total_income']:,.2f}",
            f"Rs. {summary['total_expenses']:,.2f}",
            f"Rs. {summary['net_savings']:,.2f}",
            f"{summary['savings_rate']}%",
            f"Rs. {summary['total_balance']:,.2f}",
        ],
    ]
    summary_table = Table(summary_rows, colWidths=[100, 105, 100, 90, 120])
    summary_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F1F5F9")),
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
        ("TEXTCOLOR", (0, 1), (0, 1), colors.HexColor("#10B981")),
        ("TEXTCOLOR", (1, 1), (1, 1), colors.HexColor("#EF4444")),
        ("TEXTCOLOR", (2, 1), (2, 1), colors.HexColor("#10B981")),
        ("TEXTCOLOR", (3, 1), (3, 1), colors.HexColor("#8B5CF6")),
        ("TEXTCOLOR", (4, 1), (4, 1), colors.HexColor("#6366F1")),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ("TOPPADDING", (0, 0), (-1, -1), 10),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
    ]))
    story.append(summary_table)

    story.append(Paragraph("Spending by Category", s["heading"]))

    if category_rows:
        cat_total = sum(r["total"] for r in category_rows)
        cat_table_rows = [["Category", "Total (INR)", "% of Total"]]
        for r in category_rows:
            pct = round((r["total"] / cat_total) * 100, 1) if cat_total > 0 else 0
            cat_table_rows.append([r["category"], f"Rs. {r['total']:,.2f}", f"{pct}%"])
        cat_table_rows.append(["TOTAL", f"Rs. {cat_total:,.2f}", "100%"])

        cat_table = Table(cat_table_rows, colWidths=[200, 150, 100], repeatRows=1)
        cat_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#8B5CF6")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
            ("ALIGN", (1, 1), (2, -1), "RIGHT"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F8FAFC")]),
            ("TOPPADDING", (0, 0), (-1, -1), 7),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
        ]))
        story.append(cat_table)
    else:
        story.append(Paragraph("No expenses recorded for this period.", s["empty"]))

    story.append(Paragraph(f"Monthly Trend (Last {months} Months)", s["heading"]))

    if trend_points:
        trend_table_rows = [["Month", "Income (INR)", "Expenses (INR)", "Net (INR)"]]
        for p in trend_points:
            net = p["total_income"] - p["total_expenses"]
            trend_table_rows.append([
                p["label"],
                f"Rs. {p['total_income']:,.2f}",
                f"Rs. {p['total_expenses']:,.2f}",
                f"Rs. {net:,.2f}",
            ])

        trend_table = Table(trend_table_rows, colWidths=[120, 130, 130, 120], repeatRows=1)
        trend_style_commands = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#10B981")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ALIGN", (1, 1), (3, -1), "RIGHT"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
            ("TOPPADDING", (0, 0), (-1, -1), 7),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
        ]
        for row_index, p in enumerate(trend_points, start=1):
            net = p["total_income"] - p["total_expenses"]
            net_color = colors.HexColor("#EF4444") if net < 0 else colors.HexColor("#10B981")
            trend_style_commands.append(("TEXTCOLOR", (3, row_index), (3, row_index), net_color))
            trend_style_commands.append(("FONTNAME", (3, row_index), (3, row_index), "Helvetica-Bold"))
        trend_table.setStyle(TableStyle(trend_style_commands))
        story.append(trend_table)
    else:
        story.append(Paragraph("No trend data available.", s["empty"]))

    story.append(Paragraph("Savings Goals Progress", s["heading"]))

    if goal_rows:
        goal_table_rows = [["Goal", "Target", "Saved", "Progress", "Target Date", "Status"]]
        for g in goal_rows:
            goal_table_rows.append([
                g["title"],
                f"Rs. {g['target_amount']:,.2f}",
                f"Rs. {g['current_amount']:,.2f}",
                f"{g['progress_percent']}%",
                g["target_date"].strftime("%Y-%m-%d") if g["target_date"] else "—",
                g["status"].replace("_", " ").title(),
            ])

        goal_table = Table(goal_table_rows, colWidths=[115, 80, 80, 60, 75, 75], repeatRows=1)
        goal_style_commands = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F59E0B")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ALIGN", (1, 1), (3, -1), "RIGHT"),
            ("ALIGN", (4, 1), (5, -1), "CENTER"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
            ("TOPPADDING", (0, 0), (-1, -1), 7),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ]
        for row_index, g in enumerate(goal_rows, start=1):
            status_color = colors.HexColor("#10B981") if g["progress_percent"] >= 100 else colors.HexColor("#6366F1")
            goal_style_commands.append(("TEXTCOLOR", (5, row_index), (5, row_index), status_color))
            goal_style_commands.append(("FONTNAME", (5, row_index), (5, row_index), "Helvetica-Bold"))
        goal_table.setStyle(TableStyle(goal_style_commands))
        story.append(goal_table)
    else:
        story.append(Paragraph("No savings goals set up yet.", s["empty"]))

    story.append(Spacer(1, 20))
    story.append(Paragraph("All amounts are in Indian Rupees (Rs.).", s["footer"]))

    document.build(story)
    buffer.seek(0)

    filename = f"budget-buddy-analytics-report-{year}-{month:02d}.pdf"

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ============================================================
# GET /analytics/export/excel
# CHANGED: Premium+ only.
# ============================================================

@router.get("/export/excel")
def export_analytics_excel(
    month: int = Query(..., ge=1, le=12),
    year: int = Query(..., ge=2000, le=2100),
    months: int = Query(6, ge=1, le=12),
    db: Session = Depends(get_db),
    current_user=Depends(require_export_access),  # was: get_current_user
):
    summary = get_analytics_summary_data(month, year, db, current_user)
    category_rows = get_category_spending(month, year, db, current_user)
    trend_points = get_monthly_trend_points(months, db, current_user)
    goal_rows = get_savings_progress_rows(db, current_user)

    month_name = date(year, month, 1).strftime("%B %Y")
    generated_at = datetime.now().strftime("%d %b %Y, %I:%M %p")

    workbook = Workbook()

    summary_sheet = workbook.active
    summary_sheet.title = "Analytics Summary"

    summary_sheet["A1"] = "BUDGET BUDDY"
    summary_sheet["A1"].font = Font(bold=True, size=20)

    summary_sheet["A2"] = "ANALYTICS REPORT"
    summary_sheet["A2"].font = Font(bold=True, size=14)

    summary_sheet["A3"] = month_name
    summary_sheet["A3"].font = Font(bold=True, size=12)

    summary_sheet["A4"] = f"Generated on {generated_at}"
    summary_sheet["A4"].font = Font(italic=True, size=9, color="94A3B8")

    summary_fields = [
        ("Total Income", summary["total_income"]),
        ("Total Expenses", summary["total_expenses"]),
        ("Net Savings", summary["net_savings"]),
        ("Savings Rate (%)", summary["savings_rate"]),
        ("Total Balance", summary["total_balance"]),
    ]
    for i, (label, value) in enumerate(summary_fields, start=6):
        summary_sheet[f"A{i}"] = label
        summary_sheet[f"A{i}"].font = Font(bold=True)
        summary_sheet[f"B{i}"] = value
        summary_sheet[f"B{i}"].font = Font(bold=True)
        if "Rate" not in label:
            summary_sheet[f"B{i}"].number_format = '₹#,##0.00'

    summary_sheet["A12"] = (
        "See 'Spending by Category', 'Monthly Trend' and 'Savings Goals' "
        "tabs for the full breakdown."
    )
    summary_sheet["A12"].font = Font(italic=True, size=10, color="64748B")

    summary_sheet.column_dimensions["A"].width = 28
    summary_sheet.column_dimensions["B"].width = 22

    category_sheet = workbook.create_sheet("Spending by Category")
    category_sheet.append(["Category", "Total (INR)", "% of Total"])

    cat_total = sum(r["total"] for r in category_rows)
    for r in category_rows:
        pct = round((r["total"] / cat_total) * 100, 1) if cat_total > 0 else 0
        category_sheet.append([r["category"], r["total"], pct])
    category_sheet.append(["TOTAL", cat_total, 100])

    trend_sheet = workbook.create_sheet("Monthly Trend")
    trend_sheet.append(["Month", "Income (INR)", "Expenses (INR)", "Net (INR)"])

    for p in trend_points:
        net = p["total_income"] - p["total_expenses"]
        trend_sheet.append([p["label"], p["total_income"], p["total_expenses"], net])

    goals_sheet = workbook.create_sheet("Savings Goals")
    goals_sheet.append(["Goal", "Target", "Saved", "Progress (%)", "Target Date", "Status"])

    for g in goal_rows:
        goals_sheet.append([
            g["title"],
            g["target_amount"],
            g["current_amount"],
            g["progress_percent"],
            g["target_date"].strftime("%Y-%m-%d") if g["target_date"] else "—",
            g["status"].replace("_", " ").title(),
        ])

    header_fill = PatternFill("solid", fgColor="071A3D")
    header_font = Font(bold=True, color="FFFFFF")
    total_font = Font(bold=True)
    border = Border(bottom=Side(style="thin", color="CBD5E1"))

    money_sheets = {
        category_sheet: ["B"],
        trend_sheet: ["B", "C", "D"],
        goals_sheet: ["B", "C"],
    }

    for sheet in [category_sheet, trend_sheet, goals_sheet]:
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for row in sheet.iter_rows():
            for cell in row:
                cell.border = border

        for col_letter in money_sheets.get(sheet, []):
            for cell in sheet[col_letter][1:]:
                cell.number_format = '₹#,##0.00'

        for cell in sheet[sheet.max_row]:
            cell.font = total_font

        sheet.freeze_panes = "A2"

    column_widths = {
        category_sheet: {"A": 24, "B": 18, "C": 14},
        trend_sheet: {"A": 16, "B": 18, "C": 18, "D": 18},
        goals_sheet: {"A": 26, "B": 16, "C": 16, "D": 14, "E": 16, "F": 16},
    }
    for sheet, widths in column_widths.items():
        for col_letter, width in widths.items():
            sheet.column_dimensions[col_letter].width = width

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    filename = f"budget-buddy-analytics-report-{year}-{month:02d}.xlsx"

    return StreamingResponse(
        buffer,
        media_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )