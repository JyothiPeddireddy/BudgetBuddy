from io import BytesIO
from datetime import date, datetime
from calendar import monthrange
from typing import Optional

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, joinedload

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.database import get_db
from app.models.expense import Expense
from app.models.income import Income
from app.models.account import Account
from app.models.budget import Budget
from app.models.savings_goal import SavingsGoal
from app.core.deps import get_current_user
from app.core.permissions import has_min_role, UserRole

router = APIRouter(
    prefix="/reports",
    tags=["Reports"],
)


def bank_label(item):
    """Returns the account's bank name, falling back to account name, then a placeholder."""
    account = getattr(item, "account", None)
    if not account:
        return "—"
    return account.bank_name or account.account_name or "—"


def get_range_data(
    start_date: date,
    end_date: date,
    db: Session,
    current_user,
):
    """
    Same as get_month_data but over an arbitrary date range -- used for
    both single-month exports (range = that month's start/end) and
    "All time" exports (range = earliest possible date to today).
    """
    expenses = (
        db.query(Expense)
        .options(joinedload(Expense.account))
        .filter(
            Expense.user_id == current_user.id,
            Expense.date >= start_date,
            Expense.date <= end_date,
        )
        .order_by(Expense.date.asc())
        .all()
    )

    incomes = (
        db.query(Income)
        .options(joinedload(Income.account))
        .filter(
            Income.user_id == current_user.id,
            Income.date >= start_date,
            Income.date <= end_date,
        )
        .order_by(Income.date.asc())
        .all()
    )

    total_income = sum(float(item.amount or 0) for item in incomes)
    total_expenses = sum(float(item.amount or 0) for item in expenses)
    net_savings = total_income - total_expenses

    return (
        incomes,
        expenses,
        total_income,
        total_expenses,
        net_savings,
    )


def get_month_data(
    month: int,
    year: int,
    db: Session,
    current_user,
):
    """Kept for backward compatibility -- delegates to get_range_data for one month."""
    start_date = date(year, month, 1)
    last_day = monthrange(year, month)[1]
    end_date = date(year, month, last_day)

    incomes, expenses, total_income, total_expenses, net_savings = get_range_data(
        start_date, end_date, db, current_user
    )

    return (
        incomes,
        expenses,
        total_income,
        total_expenses,
        net_savings,
        start_date,
    )


def get_accounts_data(db: Session, current_user):
    """All of the user's accounts, plus the combined balance across them."""
    accounts = (
        db.query(Account)
        .filter(Account.user_id == current_user.id)
        .order_by(Account.account_name.asc())
        .all()
    )

    total_balance = sum(float(a.balance or 0) for a in accounts)

    return accounts, total_balance


def get_budgets_data(month: int, year: int, db: Session, current_user, expenses):
    """
    Budgets set for the selected month (month_year = "YYYY-MM"), each paired
    with how much was actually spent in that category during the same month.
    Only meaningful for a single-month export -- skipped entirely for
    "All time" exports since budgets are inherently monthly.
    """
    month_year = f"{year:04d}-{month:02d}"

    budgets = (
        db.query(Budget)
        .filter(
            Budget.user_id == current_user.id,
            Budget.month_year == month_year,
        )
        .order_by(Budget.category.asc())
        .all()
    )

    spent_by_category = {}
    for e in expenses:
        spent_by_category[e.category] = spent_by_category.get(e.category, 0) + float(e.amount or 0)

    rows = []
    for b in budgets:
        limit = float(b.monthly_limit or 0)
        spent = spent_by_category.get(b.category, 0)
        remaining = limit - spent
        status = "Over Budget" if spent > limit else "On Track"
        rows.append({
            "category": b.category,
            "limit": limit,
            "spent": spent,
            "remaining": remaining,
            "status": status,
        })

    return rows


def get_goals_data(db: Session, current_user):
    """All of the user's savings goals, with computed progress percentage."""
    goals = (
        db.query(SavingsGoal)
        .filter(SavingsGoal.user_id == current_user.id)
        .order_by(SavingsGoal.created_at.desc())
        .all()
    )

    rows = []
    for g in goals:
        target = float(g.target_amount or 0)
        current = float(g.current_amount or 0)
        progress = round((current / target) * 100) if target > 0 else 0
        rows.append({
            "title": g.title,
            "target": target,
            "current": current,
            "progress": progress,
            "target_date": g.target_date,
            "status": (g.status or "in_progress").replace("_", " ").title(),
        })

    return rows


def resolve_export_range(month: Optional[int], year: Optional[int], all_time: bool, db: Session, current_user):
    """
    Decides what date range + label to use for an export:
    - all_time=True  -> earliest transaction date (income or expense) through today
    - otherwise      -> the specific month/year given
    Returns (start_date, end_date, label_str, is_all_time).
    """
    if all_time:
        earliest_expense = (
            db.query(Expense.date)
            .filter(Expense.user_id == current_user.id)
            .order_by(Expense.date.asc())
            .first()
        )
        earliest_income = (
            db.query(Income.date)
            .filter(Income.user_id == current_user.id)
            .order_by(Income.date.asc())
            .first()
        )

        candidates = [d[0] for d in (earliest_expense, earliest_income) if d]
        start_date = min(candidates) if candidates else date.today()
        end_date = date.today()

        return start_date, end_date, "All Time", True

    # Single-month export
    start_date = date(year, month, 1)
    last_day = monthrange(year, month)[1]
    end_date = date(year, month, last_day)
    label = start_date.strftime("%B %Y")

    return start_date, end_date, label, False


# ============================================================
# PDF FINANCIAL REPORT
# ============================================================

@router.get("/export/pdf")
def export_pdf(
    month: Optional[int] = Query(None, ge=1, le=12),
    year: Optional[int] = Query(None, ge=2000, le=2100),
    all_time: bool = Query(False),  # <-- NEW
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user),
):
    is_premium_plus = has_min_role(current_user, UserRole.PREMIUM)

    start_date, end_date, month_name, is_all_time = resolve_export_range(
        month, year, all_time, db, current_user
    )

    incomes, expenses, total_income, total_expenses, net_savings = get_range_data(
        start_date, end_date, db, current_user
    )

    accounts, total_balance = get_accounts_data(db, current_user)

    # Budgets are inherently monthly -- skip entirely for All Time exports.
    budget_rows = (
        get_budgets_data(month, year, db, current_user, expenses)
        if is_premium_plus and not is_all_time
        else []
    )
    goal_rows = get_goals_data(db, current_user) if is_premium_plus else []

    generated_at = datetime.now().strftime("%d %b %Y, %I:%M %p")

    buffer = BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=40,
        rightMargin=40,
        topMargin=40,
        bottomMargin=40,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "ReportTitle", parent=styles["Title"], fontSize=20, leading=24,
        alignment=TA_CENTER, textColor=colors.HexColor("#071A3D"), spaceAfter=5,
    )
    subtitle_style = ParagraphStyle(
        "ReportSubtitle", parent=styles["Normal"], fontSize=11,
        alignment=TA_CENTER, textColor=colors.HexColor("#64748B"), spaceAfter=8,
    )
    month_style = ParagraphStyle(
        "ReportMonth", parent=styles["Normal"], fontSize=13,
        alignment=TA_CENTER, textColor=colors.HexColor("#10B981"), spaceAfter=4,
    )
    generated_style = ParagraphStyle(
        "GeneratedAt", parent=styles["Normal"], fontSize=9,
        alignment=TA_CENTER, textColor=colors.HexColor("#94A3B8"), spaceAfter=20,
    )
    heading_style = ParagraphStyle(
        "ReportHeading", parent=styles["Heading2"], fontSize=13, leading=16,
        textColor=colors.HexColor("#071A3D"), spaceBefore=16, spaceAfter=8,
    )
    empty_style = ParagraphStyle(
        "EmptySection", parent=styles["Normal"], fontSize=9.5,
        textColor=colors.HexColor("#64748B"), spaceAfter=4,
    )

    story = []

    # --------------------------------------------------------
    # HEADER
    # --------------------------------------------------------

    story.append(Paragraph("BUDGET BUDDY", title_style))
    story.append(Paragraph("FINANCIAL REPORT", subtitle_style))
    story.append(Paragraph(month_name, month_style))
    if is_all_time:
        story.append(Paragraph(
            f"Covers {start_date.strftime('%d %b %Y')} to {end_date.strftime('%d %b %Y')}",
            ParagraphStyle("RangeNote", parent=styles["Normal"], fontSize=9,
                           alignment=TA_CENTER, textColor=colors.HexColor("#64748B"), spaceAfter=4),
        ))
    story.append(Paragraph(f"Generated on {generated_at}", generated_style))

    # --------------------------------------------------------
    # SUMMARY
    # --------------------------------------------------------

    story.append(Paragraph("Financial Summary", heading_style))

    summary_data = [
        ["Total Income", "Total Expenses", "Net Savings"],
        [
            f"Rs. {total_income:,.2f}",
            f"Rs. {total_expenses:,.2f}",
            f"Rs. {net_savings:,.2f}",
        ],
    ]

    summary_table = Table(summary_data, colWidths=[170, 170, 170])
    summary_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F1F5F9")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
        ("TEXTCOLOR", (0, 1), (0, 1), colors.HexColor("#10B981")),
        ("TEXTCOLOR", (1, 1), (1, 1), colors.HexColor("#EF4444")),
        ("TEXTCOLOR", (2, 1), (2, 1), colors.HexColor("#10B981")),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ("TOPPADDING", (0, 0), (-1, -1), 10),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
    ]))

    story.append(summary_table)

    # --------------------------------------------------------
    # INCOME REPORT
    # --------------------------------------------------------

    story.append(Paragraph("Income Report", heading_style))

    income_rows = [["Date", "Source", "Bank", "Amount (INR)"]]
    for item in incomes:
        income_rows.append([
            item.date.strftime("%Y-%m-%d"),
            item.source or "Income",
            bank_label(item),
            f"Rs. {float(item.amount or 0):,.2f}",
        ])
    income_rows.append(["", "TOTAL INCOME", "", f"Rs. {total_income:,.2f}"])

    income_table = Table(income_rows, colWidths=[90, 170, 110, 120], repeatRows=1)
    income_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#10B981")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("TEXTCOLOR", (3, -1), (3, -1), colors.HexColor("#10B981")),
        ("ALIGN", (3, 1), (3, -1), "RIGHT"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F8FAFC")]),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
    ]))

    story.append(income_table)

    # --------------------------------------------------------
    # EXPENSE REPORT
    # --------------------------------------------------------

    story.append(Paragraph("Expense Report", heading_style))

    expense_rows = [["Date", "Category", "Bank", "Amount (INR)"]]
    for item in expenses:
        expense_rows.append([
            item.date.strftime("%Y-%m-%d"),
            item.category or "Miscellaneous",
            bank_label(item),
            f"Rs. {float(item.amount or 0):,.2f}",
        ])
    expense_rows.append(["", "TOTAL EXPENSES", "", f"Rs. {total_expenses:,.2f}"])

    expense_table = Table(expense_rows, colWidths=[90, 170, 110, 120], repeatRows=1)
    expense_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EF4444")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("TEXTCOLOR", (3, -1), (3, -1), colors.HexColor("#EF4444")),
        ("ALIGN", (3, 1), (3, -1), "RIGHT"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F8FAFC")]),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
    ]))

    story.append(expense_table)

    # --------------------------------------------------------
    # ACCOUNTS OVERVIEW
    # --------------------------------------------------------

    story.append(Paragraph("Accounts Overview", heading_style))

    if accounts:
        account_rows = [["Account Name", "Bank", "Type", "Balance (INR)"]]
        for a in accounts:
            account_rows.append([
                a.account_name,
                a.bank_name or "—",
                a.account_type,
                f"Rs. {float(a.balance or 0):,.2f}",
            ])
        account_rows.append(["", "", "TOTAL BALANCE", f"Rs. {total_balance:,.2f}"])

        account_table = Table(account_rows, colWidths=[150, 120, 100, 120], repeatRows=1)
        account_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#6366F1")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
            ("TEXTCOLOR", (3, -1), (3, -1), colors.HexColor("#6366F1")),
            ("ALIGN", (3, 1), (3, -1), "RIGHT"),
            ("ALIGN", (2, -1), (2, -1), "RIGHT"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F8FAFC")]),
            ("TOPPADDING", (0, 0), (-1, -1), 7),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
        ]))

        story.append(account_table)
    else:
        story.append(Paragraph("No accounts on file.", empty_style))

    # --------------------------------------------------------
    # BUDGET REPORT — Premium+, single-month only
    # --------------------------------------------------------

    if is_premium_plus and not is_all_time:
        story.append(Paragraph("Budget Report", heading_style))

        if budget_rows:
            budget_table_rows = [["Category", "Monthly Limit", "Spent", "Remaining", "Status"]]
            for b in budget_rows:
                budget_table_rows.append([
                    b["category"],
                    f"Rs. {b['limit']:,.2f}",
                    f"Rs. {b['spent']:,.2f}",
                    f"Rs. {b['remaining']:,.2f}",
                    b["status"],
                ])

            budget_table = Table(budget_table_rows, colWidths=[100, 90, 90, 90, 90], repeatRows=1)
            budget_style_commands = [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#8B5CF6")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("ALIGN", (1, 1), (3, -1), "RIGHT"),
                ("ALIGN", (4, 1), (4, -1), "CENTER"),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
            ]

            for row_index, b in enumerate(budget_rows, start=1):
                status_color = colors.HexColor("#EF4444") if b["status"] == "Over Budget" else colors.HexColor("#10B981")
                budget_style_commands.append(("TEXTCOLOR", (4, row_index), (4, row_index), status_color))
                budget_style_commands.append(("FONTNAME", (4, row_index), (4, row_index), "Helvetica-Bold"))

            budget_table.setStyle(TableStyle(budget_style_commands))
            story.append(budget_table)
        else:
            story.append(Paragraph("No budgets set for this period.", empty_style))
    elif is_premium_plus and is_all_time:
        story.append(Paragraph("Budget Report", heading_style))
        story.append(Paragraph("Budget reports are monthly and not included in All Time exports.", empty_style))

    # --------------------------------------------------------
    # SAVINGS GOALS — Premium+ only
    # --------------------------------------------------------

    if is_premium_plus:
        story.append(Paragraph("Savings Goals", heading_style))

        if goal_rows:
            goal_table_rows = [["Goal", "Target", "Saved", "Progress", "Target Date", "Status"]]
            for g in goal_rows:
                goal_table_rows.append([
                    g["title"],
                    f"Rs. {g['target']:,.2f}",
                    f"Rs. {g['current']:,.2f}",
                    f"{g['progress']}%",
                    g["target_date"].strftime("%Y-%m-%d") if g["target_date"] else "—",
                    g["status"],
                ])

            goal_table = Table(goal_table_rows, colWidths=[115, 80, 80, 60, 75, 75], repeatRows=1)
            goal_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F59E0B")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("ALIGN", (1, 1), (3, -1), "RIGHT"),
                ("ALIGN", (4, 1), (5, -1), "CENTER"),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                ("FONTSIZE", (0, 0), (-1, -1), 8.5),
            ]))

            story.append(goal_table)
        else:
            story.append(Paragraph("No savings goals set up yet.", empty_style))

    story.append(Spacer(1, 20))
    story.append(Paragraph(
        "All amounts are in Indian Rupees (Rs.).",
        ParagraphStyle("Footer", parent=styles["Normal"], fontSize=8,
                       alignment=TA_CENTER, textColor=colors.HexColor("#64748B")),
    ))

    document.build(story)
    buffer.seek(0)

    filename = (
        "budget-buddy-financial-report-all-time.pdf"
        if is_all_time
        else f"budget-buddy-financial-report-{year}-{month:02d}.pdf"
    )

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ============================================================
# EXCEL FINANCIAL REPORT
# ============================================================

@router.get("/export/excel")
def export_excel(
    month: Optional[int] = Query(None, ge=1, le=12),
    year: Optional[int] = Query(None, ge=2000, le=2100),
    all_time: bool = Query(False),  # <-- NEW
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user),
):
    is_premium_plus = has_min_role(current_user, UserRole.PREMIUM)

    start_date, end_date, month_name, is_all_time = resolve_export_range(
        month, year, all_time, db, current_user
    )

    incomes, expenses, total_income, total_expenses, net_savings = get_range_data(
        start_date, end_date, db, current_user
    )

    accounts, total_balance = get_accounts_data(db, current_user)

    budget_rows = (
        get_budgets_data(month, year, db, current_user, expenses)
        if is_premium_plus and not is_all_time
        else []
    )
    goal_rows = get_goals_data(db, current_user) if is_premium_plus else []

    generated_at = datetime.now().strftime("%d %b %Y, %I:%M %p")

    workbook = Workbook()

    # ========================================================
    # FINANCIAL SUMMARY
    # ========================================================

    summary = workbook.active
    summary.title = "Financial Summary"

    summary["A1"] = "BUDGET BUDDY"
    summary["A1"].font = Font(bold=True, size=20)

    summary["A2"] = "FINANCIAL REPORT"
    summary["A2"].font = Font(bold=True, size=14)

    summary["A3"] = month_name
    summary["A3"].font = Font(bold=True, size=12)

    if is_all_time:
        summary["A4"] = f"Covers {start_date.strftime('%d %b %Y')} to {end_date.strftime('%d %b %Y')}"
        summary["A4"].font = Font(italic=True, size=9, color="64748B")
        summary["A5"] = f"Generated on {generated_at}"
        summary["A5"].font = Font(italic=True, size=9, color="94A3B8")
        next_row = 7
    else:
        summary["A4"] = f"Generated on {generated_at}"
        summary["A4"].font = Font(italic=True, size=9, color="94A3B8")
        next_row = 6

    summary[f"A{next_row}"] = "Total Income"
    summary[f"B{next_row}"] = total_income
    summary[f"A{next_row+1}"] = "Total Expenses"
    summary[f"B{next_row+1}"] = total_expenses
    summary[f"A{next_row+2}"] = "Net Savings"
    summary[f"B{next_row+2}"] = net_savings
    summary[f"A{next_row+3}"] = "Total Account Balance"
    summary[f"B{next_row+3}"] = total_balance

    note_row = next_row + 5
    if is_all_time:
        summary[f"A{note_row}"] = (
            "Budget Report is excluded from All Time exports (budgets are monthly). "
            "See 'Income Report', 'Expense Report', 'Accounts'"
            + (", and 'Savings Goals' tabs" if is_premium_plus else " tabs")
            + " for the full itemized breakdown."
        )
    else:
        summary[f"A{note_row}"] = (
            "See 'Income Report', 'Expense Report' and 'Accounts' tabs for the "
            "full itemized breakdown."
            if not is_premium_plus else
            "See 'Income Report', 'Expense Report', 'Accounts', 'Budget Report' "
            "and 'Savings Goals' tabs for the full itemized breakdown."
        )
    summary[f"A{note_row}"].font = Font(italic=True, size=10, color="64748B")

    for row in range(next_row, next_row + 4):
        summary[f"A{row}"].font = Font(bold=True)
        summary[f"B{row}"].font = Font(bold=True)
        summary[f"B{row}"].number_format = '₹#,##0.00'

    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 25

    # ========================================================
    # INCOME REPORT
    # ========================================================

    income_sheet = workbook.create_sheet("Income Report")
    income_sheet.append(["Date", "Source", "Bank", "Amount (INR)"])
    for item in incomes:
        income_sheet.append([item.date, item.source or "Income", bank_label(item), float(item.amount or 0)])
    income_sheet.append(["", "TOTAL INCOME", "", total_income])

    # ========================================================
    # EXPENSE REPORT
    # ========================================================

    expense_sheet = workbook.create_sheet("Expense Report")
    expense_sheet.append(["Date", "Category", "Bank", "Amount (INR)"])
    for item in expenses:
        expense_sheet.append([item.date, item.category or "Miscellaneous", bank_label(item), float(item.amount or 0)])
    expense_sheet.append(["", "TOTAL EXPENSES", "", total_expenses])

    # ========================================================
    # ACCOUNTS
    # ========================================================

    accounts_sheet = workbook.create_sheet("Accounts")
    accounts_sheet.append(["Account Name", "Bank", "Type", "Balance (INR)"])
    for a in accounts:
        accounts_sheet.append([a.account_name, a.bank_name or "—", a.account_type, float(a.balance or 0)])
    accounts_sheet.append(["", "", "TOTAL BALANCE", total_balance])

    # ========================================================
    # BUDGET REPORT + SAVINGS GOALS
    # ========================================================

    all_sheets = [income_sheet, expense_sheet, accounts_sheet]
    money_sheets = {
        income_sheet: ["D"],
        expense_sheet: ["D"],
        accounts_sheet: ["D"],
    }

    if is_premium_plus and not is_all_time:
        budget_sheet = workbook.create_sheet("Budget Report")
        budget_sheet.append(["Category", "Monthly Limit", "Spent", "Remaining", "Status"])
        for b in budget_rows:
            budget_sheet.append([b["category"], b["limit"], b["spent"], b["remaining"], b["status"]])
        all_sheets.append(budget_sheet)
        money_sheets[budget_sheet] = ["B", "C", "D"]

    if is_premium_plus:
        goals_sheet = workbook.create_sheet("Savings Goals")
        goals_sheet.append(["Goal", "Target", "Saved", "Progress (%)", "Target Date", "Status"])
        for g in goal_rows:
            goals_sheet.append([
                g["title"], g["target"], g["current"], g["progress"],
                g["target_date"].strftime("%Y-%m-%d") if g["target_date"] else "—",
                g["status"],
            ])
        all_sheets.append(goals_sheet)
        money_sheets[goals_sheet] = ["B", "C"]

    # ========================================================
    # EXCEL FORMATTING
    # ========================================================

    header_fill = PatternFill("solid", fgColor="071A3D")
    header_font = Font(bold=True, color="FFFFFF")
    total_font = Font(bold=True)
    border = Border(bottom=Side(style="thin", color="CBD5E1"))

    for sheet in all_sheets:
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for row in sheet.iter_rows():
            for cell in row:
                cell.border = border

        for col_letter in money_sheets.get(sheet, []):
            for cell in sheet[col_letter][1:]:
                cell.number_format = '₹#,##0.00'

        for cell in sheet[sheet.max_row]:
            cell.font = total_font

        sheet.freeze_panes = "A2"

    column_widths = {
        income_sheet: {"A": 15, "B": 28, "C": 18, "D": 20},
        expense_sheet: {"A": 15, "B": 28, "C": 18, "D": 20},
        accounts_sheet: {"A": 26, "B": 20, "C": 18, "D": 20},
    }
    if is_premium_plus and not is_all_time:
        column_widths[budget_sheet] = {"A": 22, "B": 18, "C": 18, "D": 18, "E": 16}
    if is_premium_plus:
        column_widths[goals_sheet] = {"A": 26, "B": 16, "C": 16, "D": 14, "E": 16, "F": 16}

    for sheet, widths in column_widths.items():
        for col_letter, width in widths.items():
            sheet.column_dimensions[col_letter].width = width

    # ========================================================
    # CREATE EXCEL FILE
    # ========================================================

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    filename = (
        "budget-buddy-financial-report-all-time.xlsx"
        if is_all_time
        else f"budget-buddy-financial-report-{year}-{month:02d}.xlsx"
    )

    return StreamingResponse(
        buffer,
        media_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )